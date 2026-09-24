
"""Tests of the upload fixes from the 2026-09-23 verification round.

An .xlsx of semicolon lines whose names hold a comma or whose ISIN
fields hold a placeholder, a plain .xlsx whose cells hold a semicolon,
first data rows named like labels, headers with a label outside the
word list or one that looks like data, a title row above the header,
and cp1250 text with a byte cp1250 leaves undefined. Most tests call
parse_upload directly; the route tests fake the lookup as
tests/test_app.py does, so nothing reaches GLEIF.
"""

import csv
import io

import pytest
from openpyxl import Workbook

from app import app as flask_app
from main import routes as app_module
from core import storage
from core.models import InputError, LookupResult
from core.upload import MAX_ENTITIES, parse_upload

GOOD_ROW = ["Alfa a.s.", "CZ0005112300", "CZ", "Praha", "Ulice 1", "110 00"]
LONG_ISIN = "CZ0005112300 (kmenova akcie)"


class _FakeGleifClient:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _fake_lookup(entity, client):
    return LookupResult(notes="No LEI found in the GLEIF database."), []


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(app_module, "lookup_entity", _fake_lookup)
    monkeypatch.setattr(app_module, "GleifClient", _FakeGleifClient)
    flask_app.config["TESTING"] = True
    return flask_app.test_client()


def _xlsx(rows):
    """An .xlsx with each row on its own worksheet row ([] is blank)."""
    workbook = Workbook()
    sheet = workbook.active
    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=row_number, column=column, value=value)
    content = io.BytesIO()
    workbook.save(content)
    return content.getvalue()


def _opened_with_commas(lines):
    """An .xlsx of text lines as Excel opens them with "," as delimiter.

    Each line is split at its commas, a quote only counting at the
    start of a value, so a semicolon line lands in column A and runs
    on into the next columns wherever a value holds a comma.
    """
    return _xlsx([next(csv.reader([line])) for line in lines])


def _csv(rows, delimiter=";"):
    return "".join(delimiter.join(row) + "\n" for row in rows).encode()


def _fields(content, filename="in.xlsx"):
    return [
        (entity.name, entity.isin, entity.country, entity.town,
         entity.street, entity.zip_code)
        for entity in parse_upload(filename, content)
    ]


def _names(content, filename="in.csv"):
    return [entity.name for entity in parse_upload(filename, content)]


def _refusal(content, filename="in.csv"):
    """The (English, Czech) message parse_upload refuses it with."""
    with pytest.raises(InputError) as caught:
        parse_upload(filename, content)
    return str(caught.value), caught.value.message_cs


# Semicolon lines in an .xlsx whose names hold a comma.

def _quoted(line, quote):
    """The line with every field after the name in double quotes."""
    if not quote:
        return line
    name, *rest = line.split(";")
    return ";".join([name] + [f'"{field}"' for field in rest])


TWO_COLUMN_LINES = [
    "ČEZ, a. s.;CZ0005112300",
    "Komerční banka, a.s.;CZ0008019106",
    "Moneta Money Bank a.s.;CZ0008040318",
]


@pytest.mark.parametrize("header", [None, "Název;ISIN"])
@pytest.mark.parametrize("quote", [False, True], ids=["plain", "quoted"])
def test_two_column_lines_with_commas_in_names(header, quote):
    lines = [_quoted(line, quote) for line in TWO_COLUMN_LINES]
    content = _opened_with_commas(([header] if header else []) + lines)
    assert [row[:2] for row in _fields(content)] == [
        ("ČEZ, a. s.", "CZ0005112300"),
        ("Komerční banka, a.s.", "CZ0008019106"),
        ("Moneta Money Bank a.s.", "CZ0008040318"),
    ]


def test_lines_whose_every_name_holds_a_comma():
    # No row has a semicolon in column A at all.
    content = _opened_with_commas(TWO_COLUMN_LINES[:2])
    assert [row[:2] for row in _fields(content)] == [
        ("ČEZ, a. s.", "CZ0005112300"),
        ("Komerční banka, a.s.", "CZ0008019106"),
    ]


SIX_COLUMN_LINES = [
    "ČEZ, a. s.;CZ0005112300;CZ;Praha;Duhová 2/1444;140 53",
    "Komerční banka, a.s.;CZ0008019106;CZ;Praha;"
    "Na Příkopě 33, čp. 969;114 07",
    "Raiffeisenbank a.s.;;CZ;Praha;Hvězdova 1716/2b;140 78",
]

SIX_COLUMN_ENTITIES = [
    ("ČEZ, a. s.", "CZ0005112300", "CZ", "Praha", "Duhová 2/1444",
     "140 53"),
    ("Komerční banka, a.s.", "CZ0008019106", "CZ", "Praha",
     "Na Příkopě 33, čp. 969", "114 07"),
    ("Raiffeisenbank a.s.", None, "CZ", "Praha", "Hvězdova 1716/2b",
     "140 78"),
]


@pytest.mark.parametrize("header", [
    None,
    "Název;ISIN;Země;Město;Ulice;PSČ",
    'PARTY_FULL_NAME;"ISIN_IDENT";"COUNTRY_CODE";"ADDR_CITY_NAME";'
    '"ADDR_STREET_NAME";"ADDR_ZIP_CODE"',
])
@pytest.mark.parametrize("quote", [False, True], ids=["plain", "quoted"])
def test_six_column_lines_with_commas_in_names(header, quote):
    lines = [_quoted(line, quote) for line in SIX_COLUMN_LINES]
    content = _opened_with_commas(([header] if header else []) + lines)
    assert _fields(content) == SIX_COLUMN_ENTITIES


def test_lines_with_every_field_quoted_keep_the_comma_in_the_name():
    content = _opened_with_commas([
        '"Název";"ISIN";"Země"',
        '"ČEZ, a. s.";"CZ0005112300";"CZ"',
        '"Moneta Money Bank, a.s.";"CZ0008040318";"CZ"',
    ])
    assert [row[:3] for row in _fields(content)] == [
        ("ČEZ, a. s.", "CZ0005112300", "CZ"),
        ("Moneta Money Bank, a.s.", "CZ0008040318", "CZ"),
    ]


def test_lines_with_a_few_odd_isin_fields_are_still_lines():
    content = _opened_with_commas([
        "ČEZ, a. s.;CZ0005112300;CZ",
        "Komerční banka, a.s.;n/a;CZ",
        "Moneta Money Bank a.s.;;CZ",
    ])
    assert [row[:3] for row in _fields(content)] == [
        ("ČEZ, a. s.", "CZ0005112300", "CZ"),
        ("Komerční banka, a.s.", "n/a", "CZ"),
        ("Moneta Money Bank a.s.", None, "CZ"),
    ]


@pytest.mark.parametrize("lines, expected", [
    (["Alfa a.s.;N/A;CZ", "Beta a.s.;N/A;CZ", "Gama a.s.;CZ0005112300;CZ"],
     [("Alfa a.s.", "N/A", "CZ"), ("Beta a.s.", "N/A", "CZ"),
      ("Gama a.s.", "CZ0005112300", "CZ")]),
    (["Alfa a.s.;CZ0005112300 (akcie);CZ",
      "Beta a.s.;CZ0008019106 (akcie);CZ"],
     [("Alfa a.s.", "CZ0005112300 (akcie)", "CZ"),
      ("Beta a.s.", "CZ0008019106 (akcie)", "CZ")]),
    (["Název;ISIN;Země", "ČEZ, a. s.;-;CZ", "Komerční banka, a.s.;-;CZ",
      "Moneta Money Bank a.s.;CZ0008040318;CZ"],
     [("ČEZ, a. s.", "-", "CZ"), ("Komerční banka, a.s.", "-", "CZ"),
      ("Moneta Money Bank a.s.", "CZ0008040318", "CZ")]),
    (["Název;ISIN;Země", "ČEZ, a. s.;-;CZ"], [("ČEZ, a. s.", "-", "CZ")]),
    (["Alfa a.s.;-;CZ;Praha;Ulice 1, patro 2;110 00",
      "Beta a.s.;-;CZ;Brno;Ulice 2;602 00",
      "Gama, a.s.;CZ0005112300;CZ;Brno;Ulice 3;602 00"],
     [("Alfa a.s.", "-", "CZ"), ("Beta a.s.", "-", "CZ"),
      ("Gama, a.s.", "CZ0005112300", "CZ")]),
], ids=["n/a", "isin with a note", "comma names", "header and one row",
        "six fields"])
def test_lines_whose_isin_fields_hold_placeholders_are_lines(
    lines, expected
):
    content = _opened_with_commas(lines)
    assert [row[:3] for row in _fields(content)] == expected


def test_comma_lines_keep_their_worksheet_row_numbers():
    content = _opened_with_commas([
        "Název;ISIN", "ČEZ, a. s.;CZ0005112300", "",
        f"Komerční banka, a.s.;{LONG_ISIN}",
    ])
    english, czech = _refusal(content, "in.xlsx")
    assert english == "Row 4: the ISIN is longer than 20 characters."
    assert czech == "Řádek 4: ISIN je delší než 20 znaků."


def test_the_route_stores_comma_names_from_semicolon_lines(client):
    content = _opened_with_commas(SIX_COLUMN_LINES)
    created = client.post(
        "/api/jobs",
        data={
            "mode": "bulk",
            "file_upload": (io.BytesIO(content), "in.xlsx"),
        },
        content_type="multipart/form-data",
    )
    assert created.status_code == 200, created.get_json()
    query = storage.get_search(created.get_json()["job_id"])["query"]
    assert [(entity["name"], entity["isin"]) for entity in query] == [
        entity[:2] for entity in SIX_COLUMN_ENTITIES
    ]


# A plain .xlsx whose names hold a semicolon.

@pytest.mark.parametrize("rows", [
    [["Alpha; Beta Holdings", None, "GB"],
     ["Apple; Inc", "US0378331005", "US"]],
    [["Apple; Inc", "US0378331005", "US"]],
    [["Alpha; Beta Holdings", None, "GB"]],
    [["Alfa; Beta a.s."]],
    [["Foo; Bar s.r.o.", None, "CZ", "Praha", "Na Příkopě 28", "110 00"]],
    [["Alpha; Beta; Gamma Holdings", None, "GB"],
     ["Delta; Epsilon; Zeta Ltd", None, "GB"]],
], ids=["two rows", "one row with isin", "one row with country",
        "one cell", "one full row", "two semicolons"])
def test_names_holding_a_semicolon_keep_their_columns(rows):
    assert _fields(_xlsx(rows)) == [
        tuple(row[index] if index < len(row) else None for index in range(6))
        for row in rows
    ]


def test_one_semicolon_line_is_still_read_as_a_line():
    content = _xlsx([["Alfa a.s.;CZ0005112300;CZ"]])
    assert [row[:3] for row in _fields(content)] == [
        ("Alfa a.s.", "CZ0005112300", "CZ"),
    ]


@pytest.mark.parametrize("rows", [
    [GOOD_ROW + ["CZ0005112300;CZ0008019106"],
     ["Beta a.s.", "CZ0008040318", "CZ", "Brno", "Ulice 2", "602 00",
      "CZ0008040318;CZ0009093209"]],
    [GOOD_ROW[:4] + ["Ulice 1;"],
     ["Beta a.s.", "CZ0008040318", "CZ", "Brno", "Ulice 2;"]],
    [["Alfa a.s.", "", "CZ", "Praha", "Ulice 1;"],
     ["Beta a.s.", "", "CZ", "Brno", "Ulice 2;"]],
], ids=["isin list past the columns", "street ending in ;",
        "no isin, street ending in ;"])
def test_a_semicolon_past_column_a_keeps_the_columns(rows):
    assert _fields(_xlsx(rows)) == [
        tuple(row[index] or None if index < len(row) else None
              for index in range(6))
        for row in rows
    ]


def test_one_line_with_a_stray_comma_in_its_name_leaves_the_lines():
    content = _opened_with_commas([
        "Alfa a.s.,;CZ0005112300;CZ",
        "Beta a.s.;CZ0008019106;CZ",
        "ČEZ, a. s.;CZ0008040318;CZ",
    ])
    assert [row[:3] for row in _fields(content)] == [
        ("Alfa a.s.,", "CZ0005112300", "CZ"),
        ("Beta a.s.", "CZ0008019106", "CZ"),
        ("ČEZ, a. s.", "CZ0008040318", "CZ"),
    ]


def test_two_isins_in_the_isin_column_are_refused_by_row():
    content = _xlsx([
        ["Alfa a.s.", "CZ0005112300; CZ0008019106"],
        ["Beta a.s.", "CZ0008040318; CZ0009093209"],
    ])
    english, czech = _refusal(content, "in.xlsx")
    assert english == (
        "Row 1: the ISIN is longer than 20 characters. "
        "Row 2: the ISIN is longer than 20 characters."
    )
    assert czech == (
        "Řádek 1: ISIN je delší než 20 znaků. "
        "Řádek 2: ISIN je delší než 20 znaků."
    )


# First data rows named like labels.

@pytest.mark.parametrize("first_row", [
    ["Party City", "", "US", "Woodcliff Lake", "25 Green Pond Rd",
     "07677"],
    ["Party City", "", "US"],
    ["Client Company", "", "CZ", "Praha"],
    ["Customer Name", "", "Česko"],
    ["Legal Entity Company", "", "Germany"],
    ["Raiffeisenbank a.s.", "ISIN not available", "CZ"],
    ["Raiffeisenbank a.s.", "ISIN not available"],
    ["Party City", "US702149105X", "US"],
])
def test_a_first_row_named_like_labels_is_kept(first_row):
    rows = [first_row, GOOD_ROW]
    assert _names(_csv(rows)) == [first_row[0], "Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == [first_row[0], "Alfa a.s."]


def test_a_name_list_whose_second_name_is_made_of_labels_is_kept():
    names = ["Alfa a.s.", "Party City", "Client Company", "Beta s.r.o."]
    assert _names(_csv([[name] for name in names])) == names
    assert _names(_xlsx([[name] for name in names]), "in.xlsx") == names


@pytest.mark.parametrize("header", [
    ["Name", "", "Země registrace", "Sídlo"],
    ["Name", "ISIN", "Country"],
    ["Name", "ISIN", "Země"],
    ["Name", "ISIN", "Stát"],
    ["Name", "ISIN", "State"],
    ["Name", "ISIN", "Country code"],
    ["Name", "ISIN", "Kód země"],
    ["Name", "ISIN", "Land"],
])
def test_country_words_in_a_header_are_labels_not_countries(header):
    rows = [header, GOOD_ROW]
    assert _names(_csv(rows)) == ["Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == ["Alfa a.s."]


# Headers with a label outside the word list, or one that looks like
# data.

OTHER_HEADERS = [
    ["Název klienta", "ISIN"],
    ["Název protistrany", "ISIN"],
    ["Name of entity", "ISIN"],
    ["Organization name", "ISIN"],
    ["Company short name", "ISIN", "Country"],
    ["Název klienta", "ISIN číslo"],
    ["Název protistrany", "", "Země"],
    ["Name", "ISIN", "Country", "City", "Address 1", "ZIP"],
    ["PARTY_FULL_NAME", "ISIN_IDENT", "COUNTRY_CODE", "ADDR_CITY_NAME",
     "ADDR_LINE_1", "ADDR_ZIP_CODE"],
    ["Name", "ISIN", "Country ISO 3166", "City"],
    ["NAME", "ISIN", "CC"],
]


@pytest.mark.parametrize("header", OTHER_HEADERS)
def test_a_header_with_other_labels_is_skipped(header):
    rows = [header, GOOD_ROW]
    assert _names(_csv(rows)) == ["Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == ["Alfa a.s."]


@pytest.mark.parametrize("header", [
    ["Name", "ISIN", "Address 1", "Address 2"],
    ["PARTY_FULL_NAME", "ISIN_IDENT", "ADDR_LINE_1", "ADDR_LINE_2"],
    ["Název", "", "Adresa 1", "Adresa 2", "Adresa 3"],
    ["Name", "ISIN", "Address line 1", "Address line 2"],
])
def test_numbered_labels_do_not_make_a_header_an_entity(header):
    rows = [header, GOOD_ROW]
    assert _names(_csv(rows)) == ["Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == ["Alfa a.s."]


@pytest.mark.parametrize("first_row", [
    ["Firma 1"],
    ["Client 1", "", "CZ"],
    ["Firma 1", "", "", "", "Ulice 1"],
    ["Company 12", "", "", "Praha", "Street 1"],
])
def test_a_numbered_name_starting_a_list_is_kept(first_row):
    rows = [first_row, ["Firma 2"]]
    assert _names(_csv(rows)) == [first_row[0], "Firma 2"]
    assert _names(_xlsx(rows), "in.xlsx") == [first_row[0], "Firma 2"]


@pytest.mark.parametrize("header", [
    OTHER_HEADERS[index] for index in (0, 1, 7, 8)
])
def test_a_full_file_under_a_header_with_other_labels_is_accepted(header):
    rows = [header] + [
        [f"Firma {index} a.s.", "", "CZ"] for index in range(MAX_ENTITIES)
    ]
    assert len(parse_upload("in.csv", _csv(rows))) == MAX_ENTITIES
    assert len(parse_upload("in.xlsx", _xlsx(rows))) == MAX_ENTITIES


def test_the_route_takes_a_full_file_under_a_client_name_header(client):
    rows = [["Název klienta", "ISIN"]] + [
        [f"Firma {index} a.s.", ""] for index in range(MAX_ENTITIES)
    ]
    content = "".join(";".join(row) + "\n" for row in rows).encode("cp1250")
    created = client.post(
        "/api/jobs",
        data={
            "mode": "bulk",
            "file_upload": (io.BytesIO(content), "klienti.csv"),
        },
        content_type="multipart/form-data",
    )
    assert created.status_code == 200, created.get_json()
    query = storage.get_search(created.get_json()["job_id"])["query"]
    assert len(query) == MAX_ENTITIES
    assert query[0]["name"] == "Firma 0 a.s."


@pytest.mark.parametrize("first_row", [
    ["Alfa a.s.", "ISIN", "CZ", "Praha", "Ulice 1", "110 00"],
    ["Raiffeisenbank a.s.", "ISIN (not available)", "CZ"],
    ["Party City", "ISIN", "US", "Woodcliff Lake", "25 Green Pond Rd"],
])
def test_a_data_row_with_an_isin_label_is_kept(first_row):
    rows = [first_row, GOOD_ROW]
    assert _names(_csv(rows)) == [first_row[0], "Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == [first_row[0], "Alfa a.s."]


# A title row above the header.

@pytest.mark.parametrize("rows", [
    [["Seznam subjektů"], ["Název", "ISIN", "Země"], GOOD_ROW],
    [["Seznam subjektů", "", ""], ["Název", "ISIN", "Země"], GOOD_ROW],
    [["List of entities 2026"], [], ["Name", "ISIN", "Country"], [],
     GOOD_ROW],
    [["Seznam subjektů"], ["Název", "", "Země"], GOOD_ROW],
    # A title made of label words is a header of one cell itself.
    [["Firmy"], ["Název", "ISIN", "Země"], GOOD_ROW],
    [["Company"], [], ["Name", "ISIN", "Country"], GOOD_ROW],
], ids=["title", "title with empty cells", "title and blank rows",
        "header without isin label", "label title", "label title and blank"])
def test_a_title_row_and_the_header_below_it_are_skipped(rows):
    assert _names(_csv(rows)) == ["Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == ["Alfa a.s."]


def test_a_one_cell_header_keeps_a_one_cell_name_below_it():
    rows = [["Company"], ["Party City"], ["Alfa a.s."]]
    assert _names(_csv(rows)) == ["Party City", "Alfa a.s."]
    assert _names(_xlsx(rows), "in.xlsx") == ["Party City", "Alfa a.s."]


def test_help_says_a_title_row_above_the_header_is_skipped(client):
    html = " ".join(client.get("/").get_data(as_text=True).split())
    assert (
        "A first row of column labels (such as Name, ISIN or Country) is "
        "skipped, and so is a one-cell title row above it; otherwise the "
        "first row is searched too."
    ) in html
    assert (
        "První řádek s popisky sloupců (např. Název, ISIN nebo Země) se "
        "přeskočí, stejně jako nadpis v jediné buňce nad ním; jinak se "
        "hledá i první řádek."
    ) in html


def test_rows_under_a_title_and_header_keep_their_numbers():
    rows = [
        ["Seznam subjektů"], [], ["Název", "ISIN", "Země"], GOOD_ROW,
        ["Beta a.s.", LONG_ISIN],
    ]
    for content, filename in ((_csv(rows), "in.csv"),
                              (_xlsx(rows), "in.xlsx")):
        english, czech = _refusal(content, filename)
        assert english == "Row 5: the ISIN is longer than 20 characters."
        assert czech == "Řádek 5: ISIN je delší než 20 znaků."


@pytest.mark.parametrize("rows", [
    [["Alfa a.s."], ["Customer name"], ["Beta s.r.o."]],
    [["Alfa a.s."], ["Beta s.r.o.", "CZ0005112300", "CZ"]],
    [["Alfa a.s."]],
], ids=["weak header", "data row", "only row"])
def test_a_single_cell_first_row_without_a_header_below_is_kept(rows):
    expected = [row[0] for row in rows]
    assert _names(_csv(rows)) == expected
    assert _names(_xlsx(rows), "in.xlsx") == expected


# cp1250 text with a byte cp1250 leaves undefined.

CZECH_NAMES = [
    "Škoda Auto a.s.",
    "ČEZ, a. s.",
    "Komerční banka, a.s.",
    "Třinecké železárny, a. s.",
    "Žďas, a.s.",
]


def _lines(names):
    return "".join(f"{name};;CZ\n" for name in names)


@pytest.mark.parametrize("byte", [b"\x81", b"\x83", b"\x88", b"\x90",
                                  b"\x98"])
def test_an_undefined_cp1250_byte_keeps_the_czech_letters(byte):
    content = (
        _lines(CZECH_NAMES).encode("cp1250") + b"Bad" + byte + b"Name;;CZ\n"
    )
    assert _names(content) == CZECH_NAMES + ["Bad�Name"]


def test_mostly_cp1250_with_a_utf8_n_caron_keeps_the_czech_letters():
    # A UTF-8 "ň" is the bytes C5 88, and cp1250 leaves 0x88 undefined.
    content = (
        _lines(CZECH_NAMES).encode("cp1250")
        + _lines(["Plzeň a.s."]).encode()
    )
    assert _names(content) == CZECH_NAMES + ["PlzeĹ� a.s."]

