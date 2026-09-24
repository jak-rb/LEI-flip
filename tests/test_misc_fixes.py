
"""Assorted fixes from the 2026-09-23 test run (round 3).

- Many simultaneous decisions, and failed-attempt counts racing them,
  queue on SQLite's write lock instead of missing their
  compare-and-swap and failing with "database is locked".
- Loading the legal forms at a cold start from many threads at once
  gives every thread (and ``normalize_name``'s cache) the right result.
- Control characters that Python reads as whitespace (VT, FF, FS-US)
  export as a space instead of gluing two words together.
- The single form's invisible-only check is ``core.models.is_blank``.
- A disabled decision button does not react to hover.

The store runs on SQLite (see conftest.py); nothing here touches the
network.
"""

import csv
import functools
import io
import re
import secrets
import sys
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import app as flask_app
from main import routes as app_module
from core import address, export, models, storage
from core.models import LookupResult

ROOT = Path(__file__).resolve().parents[1]
APPLE_ISIN = "US0378331005"

#: How many requests write to one job at once in the burst tests: well
#: past the ~80 that made SQLite time out before the fix.
SIMULTANEOUS = 150

#: The XML-illegal control characters that Python (and so the matcher)
#: reads as whitespace.
WHITESPACE_CONTROLS = ["\x0b", "\x0c", "\x1c", "\x1d", "\x1e", "\x1f"]


def _run_at_once(calls):
    """Run each call on its own thread, all released together.

    Returns:
        Each call's return value (in order) and the errors raised.
    """
    barrier = threading.Barrier(len(calls))
    returned = [None] * len(calls)
    errors = []

    def run(position):
        barrier.wait()
        try:
            returned[position] = calls[position]()
        except Exception as error:  # reported by the caller's assert
            errors.append(error)

    threads = [
        threading.Thread(target=run, args=(position,))
        for position in range(len(calls))
    ]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    return returned, errors


# ---- Simultaneous writes on SQLite ----

def _lei(index):
    """The one candidate LEI of row ``index``."""
    return f"{index:020d}"


def _store_review_job(count):
    """Store a finished job of ``count`` near-misses to validate."""
    job_id = secrets.token_hex(16)
    storage.create_search(
        job_id=job_id, mode="bulk",
        query=[{"name": f"entity {i}"} for i in range(count)],
    )
    rows = [
        {
            "input": {"name": f"entity {i}"},
            "match": {"lei": None},
            "closest": [{"lei": _lei(i)}],
        }
        for i in range(count)
    ]
    assert storage.append_results(job_id, rows, 0) is not None
    return job_id


def _confirmed(index):
    """The decision confirming row ``index``'s candidate."""
    return {"status": "confirmed", "lei": _lei(index)}


def test_simultaneous_decisions_on_sqlite_all_persist():
    assert not storage.using_postgres()
    job_id = _store_review_job(SIMULTANEOUS)

    saved, errors = _run_at_once([
        functools.partial(storage.record_decision, job_id, i, _lei(i))
        for i in range(SIMULTANEOUS)
    ])

    assert errors == []
    expected = [_confirmed(i) for i in range(SIMULTANEOUS)]
    assert saved == expected
    results = storage.get_search(job_id)["results"]
    assert [row.get("decision") for row in results] == expected


def test_failed_attempts_racing_decisions_on_sqlite_all_persist():
    job_id = _store_review_job(SIMULTANEOUS)
    # Even rows take a decision, odd rows count a failed attempt.
    calls = [
        functools.partial(storage.record_decision, job_id, i, _lei(i))
        if i % 2 == 0
        else functools.partial(storage.record_failed_attempt, job_id, i)
        for i in range(SIMULTANEOUS)
    ]

    returned, errors = _run_at_once(calls)

    assert errors == []
    search = storage.get_search(job_id)
    for i in range(SIMULTANEOUS):
        if i % 2 == 0:
            assert returned[i] == _confirmed(i)
            assert search["results"][i]["decision"] == _confirmed(i)
            assert "failed_attempts" not in search["query"][i]
        else:
            assert returned[i] == 1
            assert search["query"][i]["failed_attempts"] == 1
            assert "decision" not in search["results"][i]


def test_begin_write_sends_nothing_on_postgres():
    # psycopg opens the transaction with the first statement, and the
    # compare-and-swap handles rivals there, so nothing is sent.
    sent = []

    class _RecordingRaw:
        def execute(self, sql, params=None):
            sent.append(sql)

    conn = storage._Connection.__new__(storage._Connection)
    conn._raw = _RecordingRaw()
    conn._sqlite = False
    conn.begin_write()
    assert sent == []


# ---- The cold-start load of the legal forms ----

def test_cold_legal_form_load_from_many_threads_is_consistent(monkeypatch):
    names = [
        f"{word} {i} {form}"
        for i, (word, form) in enumerate(
            [("Firma", "AG"), ("Alfa", "s.r.o."), ("Beta", "a.s."),
             ("Gamma", "GmbH"), ("Delta", "Ltd")] * 8
        )
    ]
    expected = [address.normalize_name.__wrapped__(name) for name in names]

    # A fresh process: nothing loaded, nothing cached, and no compiled
    # pattern in re's cache, so the load takes as long as at a cold
    # start. Frequent thread switches let the other threads run while
    # the first one is still building the patterns.
    monkeypatch.setattr(address, "_LEGAL_FORMS", None)
    monkeypatch.setattr(address, "_LEGAL_FORM_PATTERNS", None)
    address.normalize_name.cache_clear()
    re.purge()
    switch_interval = sys.getswitchinterval()
    sys.setswitchinterval(1e-6)
    try:
        normalized, errors = _run_at_once([
            functools.partial(address.normalize_name, name)
            for name in names
        ])
        cached = [address.normalize_name(name) for name in names]
    finally:
        sys.setswitchinterval(switch_interval)
        address.normalize_name.cache_clear()

    assert errors == []
    assert normalized == expected
    assert cached == expected


# ---- Whitespace-type control characters in the exports ----

@pytest.mark.parametrize(
    "char", WHITESPACE_CONTROLS,
    ids=[f"U+{ord(c):04X}" for c in WHITESPACE_CONTROLS],
)
def test_whitespace_controls_export_as_a_space(char):
    name = f"Alfa{char}Beta a.s."
    # The matcher reads the character as a word break...
    assert address.normalize_name(name) == "alfa beta"
    search = {"results": [{
        "input": {"name": name, "street": f"Ulice{char}1"},
        "match": {"lei": None, "notes": f"No{char}match"},
        "closest": [],
    }]}

    csv_rows = list(csv.reader(io.StringIO(export.build_csv(search))))
    sheet = load_workbook(export.build_xlsx(search)).active
    excel_rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    # ...and so do both downloads.
    for rows in (csv_rows, excel_rows):
        row = dict(zip(export.COLUMNS, rows[1]))
        assert row["Name"] == "Alfa Beta a.s."
        assert row["Street"] == "Ulice 1"
        assert row["Notes"] == "No match"


def test_other_xml_illegal_characters_are_still_dropped():
    assert export._sanitize("A\x00B\x08C\x0bD\x0eE\x1fF\ufffe") == (
        "ABC DE F"
    )
    # A dropped character still cannot hide a formula.
    assert export._sanitize("\x01=1+1") == "'=1+1"


# ---- The single form's invisible-only rule ----

@pytest.fixture
def client(monkeypatch):
    def fake_lookup(entity, client):
        return LookupResult(notes="No LEI found in the GLEIF database."), []

    class FakeGleifClient:
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    monkeypatch.setattr(app_module, "lookup_entity", fake_lookup)
    monkeypatch.setattr(app_module, "GleifClient", FakeGleifClient)
    flask_app.config["TESTING"] = True
    return flask_app.test_client()


def test_single_form_uses_the_models_blank_rule():
    assert app_module.is_blank is models.is_blank
    assert not hasattr(app_module, "_none_if_invisible")


def test_long_invisible_name_with_an_isin_is_an_isin_only_search(client):
    # Dropped before InputEntity checks the name's 500-character cap.
    response = client.post("/api/jobs", data={
        "mode": "single", "entity_name": "\u200b" * 600, "isin": APPLE_ISIN,
    })
    assert response.status_code == 200, response.get_json()
    job_id = response.get_json()["job_id"]
    stored = storage.get_search(job_id)["query"][0]
    assert (stored["name"], stored["isin"]) == (None, APPLE_ISIN)


# ---- Disabled decision buttons ----

def test_hover_rules_skip_disabled_decision_buttons():
    css = (ROOT / "src" / "main" / "static" / "styles.css").read_text(
        encoding="utf-8"
    )
    css = re.sub(r"/\*.*?\*/", "", css, flags=re.DOTALL)
    selectors = [
        selector.strip()
        for prelude in re.findall(r"([^{};]+)\{", css)
        for selector in prelude.split(",")
    ]
    # The hovered compound selector (the last one) of every hover rule
    # a decision button (".btn candidate-confirm", ".btn validate-none")
    # can match.
    hovered = [
        selector.split()[-1] for selector in selectors
        if ":hover" in selector.split()[-1]
        and re.search(
            r"\.(?:btn|candidate-confirm|validate-none)(?![\w-])",
            selector.split()[-1],
        )
    ]
    assert len(hovered) >= 3
    for compound in hovered:
        assert ":not(:disabled)" in compound, compound

