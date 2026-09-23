
"""Exports of XML-illegal text, and the speed of name normalization.

Two defects from the 2026-09-23 test run (report #2 and #8):

- Text XML 1.0 forbids (C0 control characters other than tab, LF and
  CR, and the noncharacters U+FFFE/U+FFFF) made /download/excel answer
  500, or write a workbook that no longer opens.
- A long whitespace run in a name made the legal-form patterns of
  ``normalize_name`` backtrack quadratically, so one /run took
  minutes. A copy of the old function checks that the fix leaves every
  normalized name unchanged.

GLEIF and OpenFIGI are faked; nothing here touches the network.
"""

import csv
import io
import random
import time

import pytest
import requests
from openpyxl import load_workbook
from unidecode import unidecode

import app as app_module
from core import address, export, isin, lookup, matcher
from core.models import (
    CandidateSummary,
    GleifAddress,
    GleifCandidate,
    LookupResult,
    MatchType,
)

MATCH_LEI = "M" * 20
REVIEW_LEI = "R" * 20

#: Characters XML 1.0 forbids: the C0 controls except tab, LF and CR,
#: and the noncharacters U+FFFE and U+FFFF.
XML_ILLEGAL = [
    chr(code) for code in range(0x20) if chr(code) not in "\t\n\r"
] + ["\ufffe", "\uffff"]


class _FakeGleifClient:
    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _echo_lookup(entity, client):
    """A lookup whose GLEIF strings all repeat the searched name."""
    text = entity.name
    if "review" in text.lower():
        candidate = CandidateSummary(
            legal_name=text, lei=REVIEW_LEI, status=text, country=text,
            city=text, street=text, overall=61.0, legal_address=text,
            hq_address=text,
        )
        return LookupResult(notes=text), [candidate]
    result = LookupResult(
        lei=MATCH_LEI, lei_status=text, match_type=MatchType.FULL_MATCH,
        confidence=90, gleif_legal_name=text, gleif_legal_address=text,
        gleif_hq_address=text, gleif_legal_country=text,
        gleif_legal_city=text, gleif_legal_street=text, notes=text,
        warnings=[text],
    )
    return result, []


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(app_module, "lookup_entity", _echo_lookup)
    monkeypatch.setattr(app_module, "GleifClient", _FakeGleifClient)
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def _finished_job(client, **fields):
    """Create a single-form job, run it, and return its id."""
    created = client.post("/api/jobs", data={"mode": "single", **fields})
    assert created.status_code == 200, created.get_json()
    job_id = created.get_json()["job_id"]
    assert client.post(f"/api/jobs/{job_id}/run").get_json()["done"]
    return job_id


def _exported_rows(client, job_id):
    """Both downloads of a job, as lists of rows keyed by column."""
    excel = client.get(f"/download/excel?job={job_id}")
    assert excel.status_code == 200
    sheet = load_workbook(io.BytesIO(excel.data)).active
    excel_rows = [
        ["" if value is None else value for value in row]
        for row in sheet.iter_rows(values_only=True)
    ]

    csv_response = client.get(f"/download/csv?job={job_id}")
    assert csv_response.status_code == 200
    csv_text = csv_response.get_data(as_text=True)
    csv_rows = list(csv.reader(io.StringIO(csv_text)))

    for rows in (excel_rows, csv_rows):
        assert rows[0] == export.COLUMNS
        for row in rows:
            for value in row:
                assert not any(char in str(value) for char in XML_ILLEGAL)
    return (
        [dict(zip(export.COLUMNS, row)) for row in excel_rows[1:]],
        [dict(zip(export.COLUMNS, row)) for row in csv_rows[1:]],
    )


@pytest.mark.parametrize(
    "char", XML_ILLEGAL, ids=[f"U+{ord(c):04X}" for c in XML_ILLEGAL],
)
def test_exports_drop_xml_illegal_characters(client, char):
    # Reachable from the single form, any upload, and GLEIF's own
    # strings: the Excel download used to answer 500 (IllegalCharacter
    # Error) for a control character and write a workbook openpyxl and
    # Excel cannot open for U+FFFE/U+FFFF.
    job_id = _finished_job(
        client,
        entity_name=f"Match{char}Jedna a.s.",
        isin=f"CZ{char}0005112300",
        country=f"C{char}Z",
        city=f"Pra{char}ha",
        street=f"Ulice{char} 1",
        postal_code=f"110{char}00",
    )
    for rows in _exported_rows(client, job_id):
        row = rows[0]
        assert (
            row["Name"], row["ISIN"], row["Country"], row["Town"],
            row["Street"], row["ZIP code"],
        ) == ("MatchJedna a.s.", "CZ0005112300", "CZ", "Praha",
              "Ulice 1", "11000")
        assert row["LEI"] == MATCH_LEI
        # Every GLEIF string of the fake repeats the name.
        for column in (
            "LEI_status", "Warnings", "GLEIF_legal_name",
            "GLEIF_legal_address", "GLEIF_hq_address", "Notes",
        ):
            assert row[column] == "MatchJedna a.s.", column


@pytest.mark.parametrize(
    "char", XML_ILLEGAL, ids=[f"U+{ord(c):04X}" for c in XML_ILLEGAL],
)
def test_exports_drop_xml_illegal_characters_of_a_candidate(client, char):
    # A near-miss exports the lookup's notes; once the user confirms the
    # candidate, its GLEIF name, status and addresses are exported.
    job_id = _finished_job(client, entity_name=f"Review{char}Ltd")
    for rows in _exported_rows(client, job_id):
        assert rows[0]["Name"] == "ReviewLtd"
        assert rows[0]["Notes"] == "ReviewLtd"

    confirmed = client.post(
        "/api/decision",
        json={"job_id": job_id, "index": 0, "choice": REVIEW_LEI},
    )
    assert confirmed.status_code == 200
    for rows in _exported_rows(client, job_id):
        row = rows[0]
        assert (row["LEI"], row["Match_type"]) == (
            REVIEW_LEI, "MANUAL_MATCH",
        )
        for column in (
            "LEI_status", "GLEIF_legal_name", "GLEIF_legal_address",
            "GLEIF_hq_address",
        ):
            assert row[column] == "ReviewLtd", column


def test_exports_neutralise_a_formula_behind_a_control_character(client):
    # The characters are dropped before the formula check, so "\x01=..."
    # cannot turn into a live formula once the \x01 is gone.
    job_id = _finished_job(client, entity_name="\x01=1+1 Match AG")
    for rows in _exported_rows(client, job_id):
        assert rows[0]["Name"] == "'=1+1 Match AG"
        assert rows[0]["GLEIF_legal_name"] == "'=1+1 Match AG"


# normalize_name: the same output, without quadratic backtracking.


def _normalize_name_before_fix(name):
    """normalize_name as of 5b0b3ae, the reference output."""
    if not name:
        return ""

    result = name.strip().lower()

    address._load_legal_forms()
    for pattern in address._LEGAL_FORM_PATTERNS:
        result = pattern.sub(' ', result)

    # Strip share-class suffixes (e.g. "- BI EUR", "Class A", "(Acc)").
    result = address._RE_SHARE_CLASS_SUFFIX.sub('', result)
    result = address._RE_SHARE_CLASS_WORD.sub('', result)
    result = address._RE_TRAILING_PARENS.sub('', result)

    result = unidecode(result)

    result = address._RE_COMMA_TRAIL.sub(' ', result)
    result = address._RE_WHITESPACE.sub(' ', result)
    return result.strip()


#: Realistic names: Czech ones with diacritics, share classes,
#: parentheses, and the names used across the tests and the fixture.
SAMPLE_NAMES = [
    "Raiffeisenbank a.s.",
    "ČEZ, a. s.",
    "Komerční banka, a.s.",
    "Česká spořitelna, a.s.",
    "Škoda Auto a.s.",
    "Pražská energetika, a.s.",
    "ŘÍZENÍ LETOVÉHO PROVOZU ČESKÉ REPUBLIKY, státní podnik",
    "Investiční společnost České spořitelny, a.s.",
    "Kooperativa pojišťovna, a.s., Vienna Insurance Group",
    "MONETA Money Bank, a.s.",
    "Jihočeská plynárenská, a.s.",
    "Družstevní záložna PSD",
    "Nadace Via, z. ú.",
    "Město Třebíč",
    "AGROFERT, a.s.",
    "Česká zbrojovka a.s.",
    "Vodafone Czech Republic a.s.",
    "Nobody s.r.o.",
    "Firma s. r. o.",
    "Conseq Invest Akciový fond - třída A",
    "Amundi Funds - Pioneer US Bond - A EUR",
    "Fidelity Funds - Asia Fund - A DIS USD",
    "iShares Core MSCI World UCITS ETF USD (Acc)",
    "Vanguard FTSE All-World UCITS ETF - USD Accumulating",
    "BlackRock Global Funds - World Mining Fund Class A2",
    "Generali Investments Podílový fond (CZK)",
    "J&T Bond CZK Class A",
    "Allianz Technology SE",
    "Allianz",
    "Apex Capital",
    "Summit Capital",
    "Johnson & Johnson",
    "Coca-Cola Company",
    "Tesco Property Finance 1 PLC",
    "Tesco Property Finance 3 PLC",
    "Barclays PLC",
    "Apple Inc",
    "ACME HOLDING a.s.",
    "ACME HOLDING 7 a.s.",
    "ACME 7 GROUP s.r.o.",
    "Match AG",
    "Review Ltd",
    "CBRE Investment Management Listed Real Assets LLC",
    "Real REMAX Group Inc",
    "Abacus Global Management Inc",
    "Longeveron Inc",
    "Nomura Holdings INC",
    "FIRY INC",
    "Deutsche Bank Aktiengesellschaft",
    "Volkswagen Financial Services (UK) Limited",
    "BNP PARIBAS S.A.",
    "Société Générale S.A.",
    "UniCredit Bank Czech Republic and Slovakia, a.s.",
    "ING Bank N.V.",
    "Toyota Motor Kabushiki Kaisha",
    "Macquarie Group Pty Ltd",
    "PKO Bank Polski Sp. z o.o.",
    "Zagrebačka banka d.d.",
    "Raiffeisen Bank International AG (RBI)",
    "Fond (Holdings) Limited",
    "(Old) Name Co.",
    "Foo Bar (a very long parenthetical note)",
    "Alfa, Beta; Gama: Delta / Epsilon",
    "A.B.C. s.r.o.",
    "7-Eleven, Inc.",
]

#: A sample of what ``\s`` matches: ASCII whitespace, the no-break and
#: ideographic spaces, the line separators, and U+0085, the one that
#: unidecode drops.
WHITESPACE = [
    " ", "\t", "\n", "\r", "\x0b", "\x0c", "\x1c", "\x85", "\u00a0",
    "\u2028", "\u2029", "\u3000",
]


def _legal_form_names():
    """Every legal form attached to names with varied separators."""
    forms = (
        (address.DATA_DIR / "legal_forms.txt")
        .read_text(encoding="utf-8").split("\n")
    )
    separators = [" ", ", ", ",", " ,", "  ", "\t", "\u00a0", " - ",
                  "\n", ". "]
    bases = ["Acme", "Komerční banka", "Tesco Property Finance 3"]
    names = []
    for form in filter(None, (line.strip() for line in forms)):
        variants = {form, form.upper(), form.title()}
        if " " in form:
            # The same form with irregular spacing inside it.
            variants |= {form.replace(" ", ws) for ws in ("  ", "\t",
                                                          "\u00a0")}
        for variant in sorted(variants):
            for index, separator in enumerate(separators):
                names.append(bases[index % 3] + separator + variant)
            names += [
                f"Acme, {variant}.",
                f"{variant} Acme",
                f"Acme {variant} Holding",
                f"Acme ({variant})",
                f"ČEZ {variant}, a. s.",
            ]
    return names


def _whitespace_run_names():
    """Names with whitespace runs of every length around the cut-off."""
    templates = [
        "Acme{ws}a.s.", "Acme a.s.{ws}Holding", "ČEZ,{ws}a. s.",
        "Foo{ws}Pty Ltd", "Foo Pty{ws}Ltd", "Foo s.{ws}r. o.",
        "Fund Class A{ws}Extra", "Fund class{ws}a{ws}x",
        "Fund (Acc{ws})", "Fund ({ws}Acc)", "Fund - A{ws}USD",
        "Alpha{ws}Beta", "{ws}Acme{ws}", "Acme{ws},{ws}Ltd",
        "Acme,{ws}ltd{ws}.", "akciová{ws}společnost Foo",
        "Foo{ws}\x85{ws}Bar",
    ]
    rng = random.Random(20260923)
    names = []
    for length in (1, 2, 21, 22, 23, 45):
        runs = [char * length for char in " \t\n\x85\u00a0\u2028"]
        runs += [
            "".join(rng.choice(WHITESPACE) for _ in range(length))
            for _ in range(2)
        ]
        for template in templates:
            names += [template.format(ws=run) for run in runs]
    return names


def _fuzzed_names(count=800):
    """Random names of legal forms, punctuation and whitespace."""
    rng = random.Random(8)
    fragments = address._load_legal_forms() + [
        "(", ")", "(acc)", " class a", " - a usd", ",", ".", "-", "acme",
        "x", "1", "ČEZ", "třída", "&",
    ]
    names = []
    for _ in range(count):
        name = ""
        for _ in range(rng.randint(1, 8)):
            name += rng.choice(fragments)
            roll = rng.random()
            if roll < 0.5:
                name += " "
            elif roll < 0.8:
                length = rng.choice([1, 2, 3, 20, 21, 22, 25, 40])
                name += "".join(
                    rng.choice(WHITESPACE) for _ in range(length)
                )
            elif roll < 0.9:
                name += rng.choice(WHITESPACE) * rng.choice([21, 22, 40])
        names.append(name)
    return names


#: Names at the 500-character limit (the old code needs about 0.4 s
#: for each of these, so there are only two).
LONG_NAMES = [
    "A" + " " * 498 + "B",
    # A newline deep inside a run stops the share-class ".*".
    "Fund class a" + " " * 400 + "\n" + " " * 86 + "x",
]


def test_normalize_name_output_is_unchanged():
    # Shortening long whitespace runs must not change any result: the
    # matcher's thresholds are audit-validated on the old output.
    names = (
        SAMPLE_NAMES + _legal_form_names() + _whitespace_run_names()
        + _fuzzed_names() + LONG_NAMES
    )
    changed = []
    for name in names:
        before = _normalize_name_before_fix(name)
        after = address.normalize_name(name)
        if after != before:
            changed.append((name, before, after))
    assert changed == []


#: Hostile names at the 500-character limit.
HOSTILE_NAMES = {
    "spaces": "A" + " " * 498 + "B",
    "tabs": "A" + "\t" * 498 + "B",
    "no-break spaces": "A" + "\u00a0" * 498 + "B",
    "newlines": "A" + "\n" * 498 + "B",
    "mixed whitespace": "A" + " \t\n\u00a0\x85\u3000" * 83 + "B",
    "spaces then a legal form": "A" + " " * 493 + " ltd B",
    "share class then spaces": "Fund class a" + " " * 487 + "x",
    "parenthesis then spaces": "Fund (" + " " * 493 + "x",
    "runs of 21 spaces": (("x" + " " * 21) * 23)[:500],
    "commas": "A" + "," * 498 + "B",
    "comma space": "A" + ", " * 249 + "B",
    "space comma": "A" + " ," * 249 + "B",
    "dots": "A" + "." * 498 + "B",
    "dashes": "A" + " -" * 249 + "B",
    "parentheses": "(" * 500,
    "legal forms": ("s.r.o. a.s. ltd " * 32)[:500],
    # Removing 166 "bv" leaves a run of spaces the shorter legal-form
    # patterns then scan: bounded, as a run the loop itself creates
    # has at most one character per removed form.
    "repeated legal form": "A" + " bv" * 166 + "B",
}


def _forget_normalized_names():
    """Empty normalize_name's cache, so a timed call does the work."""
    getattr(address.normalize_name, "cache_clear", lambda: None)()


def _fast_enough(function, text, limit=0.05, attempts=10):
    """Whether an uncached call beats ``limit`` within ``attempts``."""
    # Retrying rides out a busy machine (other test runs can take every
    # core); the backtracking this guards against was slow every time.
    for _ in range(attempts):
        _forget_normalized_names()
        started = time.perf_counter()
        function(text)
        if time.perf_counter() - started < limit:
            return True
    return False


@pytest.mark.parametrize("label", list(HOSTILE_NAMES))
def test_normalizers_are_fast_on_hostile_names(label):
    # "A", 498 spaces, "B" took 0.3-1 s per normalize_name call, which
    # the matcher makes about 9 times per candidate: one /run took
    # minutes. Every call on the lookup path must stay under 50 ms.
    text = HOSTILE_NAMES[label]
    assert len(text) == 500
    functions = {
        "normalize_name": address.normalize_name,
        "normalize_address_part": address.normalize_address_part,
        "extract_zip": address.extract_zip,
        "country_to_iso": address.country_to_iso,
        "name_similarity": (
            lambda s: matcher.name_similarity(s, "ACME HOLDING a.s.")
        ),
        "name_similarity (GLEIF side)": (
            lambda s: matcher.name_similarity("ACME HOLDING a.s.", s)
        ),
        "city_similarity": lambda s: matcher.city_similarity(s, "Praha"),
        "street_similarity": (
            lambda s: matcher.street_similarity(s, "Václavské náměstí 1")
        ),
    }
    address.normalize_name("warm up")  # loads the legal forms
    slow = [
        name for name, function in functions.items()
        if not _fast_enough(function, text)
    ]
    assert slow == []


class _CandidateGleifClient:
    """A GLEIF client whose every name search finds the same records."""

    def __init__(self, candidates):
        self._candidates = candidates

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def search_by_name(self, name, country=None, page_size=10):
        return list(self._candidates)

    def search_by_name_no_country(self, name, page_size=10):
        return list(self._candidates)


def _no_network(*args, **kwargs):
    raise AssertionError("the test tried to reach the network")


def test_run_with_a_whitespace_heavy_name_is_fast(monkeypatch):
    # The crazy-test repro: 20 plausible candidates and a single-form
    # name of "A", 498 spaces, "B". One /run took 44-170 s, and on
    # Vercel it would pass the 300 s limit with 40 candidates.
    candidates = [
        GleifCandidate(
            lei=f"{index:020d}", legal_name=f"ACME HOLDING {index} a.s.",
            status="ISSUED", other_names=[f"ACME {index} GROUP s.r.o."] * 2,
            legal_address=GleifAddress(
                country="CZ", city="Praha", postal_code="11000",
                address_lines=["Václavské náměstí 1"],
            ),
        )
        for index in range(20)
    ]
    monkeypatch.setattr(
        app_module, "GleifClient",
        lambda *args, **kwargs: _CandidateGleifClient(candidates),
    )
    monkeypatch.setattr(lookup, "resolve_isin_to_names", _no_network)
    monkeypatch.setattr(isin, "resolve_isin_to_names", _no_network)
    monkeypatch.setattr(requests.Session, "request", _no_network)
    app_module.app.config["TESTING"] = True
    client = app_module.app.test_client()

    created = client.post("/api/jobs", data={
        "mode": "single", "entity_name": "A" + " " * 498 + "B",
        "country": "CZ", "city": "Praha",
    })
    job_id = created.get_json()["job_id"]
    _forget_normalized_names()  # other tests use this very name
    started = time.perf_counter()
    response = client.post(f"/api/jobs/{job_id}/run")
    elapsed = time.perf_counter() - started

    assert response.status_code == 200
    assert response.get_json()["done"] is True
    assert elapsed < 5

