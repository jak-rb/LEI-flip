
"""Tests of the bulk-upload parser in core/upload.py.

Most tests call parse_upload directly; the route test checks that its
messages reach the search page in both languages. Creating a job looks
nothing up, and the lookup is faked to fail loudly if it ever did.
"""

import io
import time
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

import app as app_module
import core.upload as upload
from core.models import InputError
from core.upload import MAX_ENTITIES, parse_upload

HEADER = ["Name", "ISIN", "Country", "City", "Street", "Postal code"]
GOOD_ROW = ["Alfa a.s.", "CZ0005112300", "CZ", "Praha", "Ulice 1", "110 00"]
LONG_ISIN = "CZ0005112300 (kmenova akcie)"
TOO_MANY_EN = f"Too many entities (more than {MAX_ENTITIES})."
TOO_MANY_CS = f"Příliš mnoho subjektů (více než {MAX_ENTITIES})."
TOO_MUCH_DATA_EN = "The .xlsx file holds too much data to read."
TOO_MUCH_DATA_CS = "Soubor .xlsx obsahuje příliš mnoho dat."


def _csv(rows, delimiter=";"):
    return "".join(delimiter.join(row) + "\n" for row in rows).encode()


def _xlsx(rows):
    """An .xlsx with each row on its own worksheet row ([] is blank)."""
    workbook = Workbook()
    sheet = workbook.active
    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            if value is not None:
                sheet.cell(row=row_number, column=column, value=value)
    content = io.BytesIO()
    workbook.save(content)
    return content.getvalue()


def _with_parts(content, parts):
    """Rewrite an .xlsx, replacing or adding the given zip members."""
    source = zipfile.ZipFile(io.BytesIO(content))
    result = io.BytesIO()
    with zipfile.ZipFile(
        result, "w", zipfile.ZIP_DEFLATED, compresslevel=1
    ) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename in parts:
                replace = parts[item.filename]
                data = replace(data) if callable(replace) else replace
            target.writestr(item.filename, data)
        for name, data in parts.items():
            if name not in source.namelist():
                target.writestr(name, data)
    return result.getvalue()


def _sheet_xml(rows_xml, dimension=None):
    dimension = f'<dimension ref="{dimension}"/>' if dimension else ""
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/'
        f'spreadsheetml/2006/main">{dimension}<sheetData>{rows_xml}'
        "</sheetData></worksheet>"
    ).encode()


def _inline_row(number, *values):
    cells = "".join(
        f'<c r="{column}{number}" t="inlineStr"><is><t>{value}</t></is></c>'
        for column, value in zip("ABCDEF", values)
    )
    return f'<row r="{number}">{cells}</row>'


def _entities(content, filename="in.csv"):
    return [
        (entity.name, entity.isin)
        for entity in parse_upload(filename, content)
    ]


def _refusal(content, filename="in.csv"):
    """The (English, Czech) message parse_upload refuses it with."""
    with pytest.raises(InputError) as caught:
        parse_upload(filename, content)
    return str(caught.value), caught.value.message_cs


@pytest.fixture
def client(monkeypatch):
    def _no_lookup(entity, client_):
        raise AssertionError("creating a job must not look anything up")
    monkeypatch.setattr(app_module, "lookup_entity", _no_lookup)
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


# Rows are never dropped silently.

@pytest.mark.parametrize(("column", "field_en", "field_cs", "limit"), [
    (0, "the name", "název", 500),
    (1, "the ISIN", "ISIN", 20),
    (2, "the country", "země", 200),
    (3, "the city", "město", 200),
    (4, "the street", "ulice", 500),
    (5, "the postal code", "PSČ", 20),
])
def test_an_over_long_value_names_its_row_field_and_limit(
    column, field_en, field_cs, limit
):
    bad_row = ["Beta a.s.", "", "", "", "", ""]
    bad_row[column] = "x" * (limit + 1)
    english, czech = _refusal(_csv([HEADER, GOOD_ROW, bad_row]))
    assert english == f"Row 3: {field_en} is longer than {limit} characters."
    assert czech == f"Řádek 3: {field_cs} je delší než {limit} znaků."


def test_a_file_whose_only_row_is_too_long_says_why():
    # It used to answer "No entities found", which was wrong.
    english, czech = _refusal(_csv([["Beta a.s.", LONG_ISIN]]))
    assert english == "Row 1: the ISIN is longer than 20 characters."
    assert czech == "Řádek 1: ISIN je delší než 20 znaků."


def test_one_row_with_two_long_values_names_both():
    row = ["Beta a.s.", LONG_ISIN, "CZ", "Praha", "", "1" * 21]
    english, czech = _refusal(_csv([row]))
    assert english == (
        "Row 1: the ISIN is longer than 20 characters; the postal code "
        "is longer than 20 characters."
    )
    assert czech == (
        "Řádek 1: ISIN je delší než 20 znaků; PSČ je delší než 20 znaků."
    )


def test_text_row_numbers_count_the_header_and_blank_lines():
    content = (
        "Name;ISIN\n\nAlfa a.s.;CZ0005112300\n\n"
        f"Beta a.s.;{LONG_ISIN}\n"
    ).encode()
    english, _ = _refusal(content)
    assert english.startswith("Row 5: ")


def test_xlsx_row_numbers_are_the_worksheet_rows():
    rows = [HEADER, GOOD_ROW, [], [], [], [], ["Beta a.s.", LONG_ISIN]]
    english, czech = _refusal(_xlsx(rows), "in.xlsx")
    assert english.startswith("Row 7: ")
    assert czech.startswith("Řádek 7: ")


def test_semicolon_lines_in_xlsx_keep_their_worksheet_row_numbers():
    rows = [
        ["Name;ISIN"], ["Alfa a.s.;CZ0005112300"], [],
        [f"Beta a.s.;{LONG_ISIN}"],
    ]
    english, _ = _refusal(_xlsx(rows), "in.xlsx")
    assert english == "Row 4: the ISIN is longer than 20 characters."


@pytest.mark.parametrize(("bad_rows", "more_en", "more_cs"), [
    (6, "And 1 more row.", "A ještě 1 další řádek."),
    (8, "And 3 more rows.", "A ještě 3 další řádky."),
    (12, "And 7 more rows.", "A ještě 7 dalších řádků."),
])
def test_at_most_five_rows_are_listed_and_the_rest_counted(
    bad_rows, more_en, more_cs
):
    rows = [[f"Firma {index} a.s.", LONG_ISIN] for index in range(bad_rows)]
    english, czech = _refusal(_csv(rows))
    listed = [f"Row {number}: " in english for number in range(1, 7)]
    assert listed == [True] * 5 + [False]
    assert english.endswith(" " + more_en)
    assert czech.endswith(" " + more_cs)


def test_an_invalid_row_still_counts_toward_the_cap():
    rows = [GOOD_ROW] * MAX_ENTITIES + [["Beta a.s.", LONG_ISIN]]
    english, czech = _refusal(_csv(rows))
    assert english.startswith(TOO_MANY_EN)
    assert czech.startswith(TOO_MANY_CS)


def test_the_route_returns_the_row_error_in_both_languages(client):
    content = _csv([HEADER, GOOD_ROW, ["Beta a.s.", LONG_ISIN]])
    response = client.post(
        "/api/jobs",
        data={
            "mode": "bulk",
            "file_upload": (io.BytesIO(content), "in.csv"),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert response.get_json() == {
        "error": "Row 3: the ISIN is longer than 20 characters.",
        "error_cs": "Řádek 3: ISIN je delší než 20 znaků.",
    }


# Header detection.

@pytest.mark.parametrize("header", [
    ["Název subjektu", "ISIN kód", "Země"],
    ["Název subjektu", "", "Země"],
    ["Entity name", "ISIN", "Country", "City", "Street", "Postal code"],
    ["Legal name", "Identifier"],
    ["PARTY_FULL_NAME", "ISIN_IDENT", "COUNTRY_CODE", "ADDR_CITY_NAME",
     "ADDR_STREET_NAME", "ADDR_ZIP_CODE"],
    ["Subjekt", "Kód", "Země", "Město", "Ulice", "PSČ"],
    ["Obchodní firma", "Kód ISIN"],
    ["Protistrana", "", "Country", "City", "Street", "ZIP"],
    ["", "ISIN"],
    ["Counterparty", "ISIN (if known)"],
    ["Name (required)", "ISIN (optional)"],
    ["Klient", "ISIN číslo"],
    ["Issuer", "ISINCODE"],
    ["Counterparty"],
    ["Client"],
    ["Customer name"],
])
def test_natural_and_database_headers_are_skipped(header):
    content = _csv([header, GOOD_ROW])
    assert _entities(content) == [("Alfa a.s.", "CZ0005112300")]
    assert _entities(_xlsx([header, GOOD_ROW]), "in.xlsx") == [
        ("Alfa a.s.", "CZ0005112300"),
    ]


@pytest.mark.parametrize("first_row", [
    ["Name", "CZ0005112300"],
    ["Company", "US0378331005", "US"],
    ["Acme a.s.", "ISIN CZ0005112300"],
    ["Country Garden Holdings Co Ltd", "", "CN"],
    ["City Developments Ltd", "", "SG"],
    ["Street Capital Group Inc"],
    # A city or street named like a label (Clarks is in Street).
    ["C. & J. Clark International Limited", "", "GB", "Street",
     "40 High Street", "BA16 0EQ"],
    ["Beta a.s.", "", "GB", "Leeds", "Town Street", "LS1 6PU"],
    ["Beta a.s.", "", "CZ", "Praha", "Ulice", "110 00"],
])
def test_a_data_row_is_never_taken_for_a_header(first_row):
    names = [name for name, _ in _entities(_csv([first_row, GOOD_ROW]))]
    assert names == [first_row[0], "Alfa a.s."]
    names = [
        name for name, _ in _entities(_xlsx([first_row, GOOD_ROW]), "in.xlsx")
    ]
    assert names == [first_row[0], "Alfa a.s."]


def test_a_header_with_a_worded_isin_label_leaves_room_for_100_rows():
    rows = [["Counterparty", "ISIN (if known)"]] + [
        [f"Firma {index} a.s.", ""] for index in range(MAX_ENTITIES)
    ]
    assert len(parse_upload("in.csv", _csv(rows))) == MAX_ENTITIES


def test_the_header_after_blank_rows_is_skipped():
    blank_first = b"\nName,ISIN,Country\nCEZ a. s.,CZ0005112300,Cesko\n"
    assert _entities(blank_first) == [("CEZ a. s.", "CZ0005112300")]

    empty_cells = (
        ";;;;;\r\nNázev;ISIN;Země\r\nČEZ, a. s.;CZ0005112300;CZ\r\n"
    ).encode("cp1250")
    assert _entities(empty_cells) == [("ČEZ, a. s.", "CZ0005112300")]

    utf16 = "\r\nNázev\tISIN\r\nČEZ, a. s.\tCZ0005112300\r\n"
    assert _entities(utf16.encode("utf-16"), "in.txt") == [
        ("ČEZ, a. s.", "CZ0005112300"),
    ]


def test_a_full_xlsx_whose_header_is_on_row_two_is_accepted():
    rows = [[], ["Name", "ISIN", "Country"]] + [
        [f"Company {index}", None, "CZ"] for index in range(MAX_ENTITIES)
    ]
    entities = parse_upload("in.xlsx", _xlsx(rows))
    assert len(entities) == MAX_ENTITIES
    assert entities[0].name == "Company 0"


def test_help_says_which_first_row_is_skipped(client):
    html = client.get("/").get_data(as_text=True)
    assert "header names don't matter" not in html
    assert "A first row of column labels" in html
    assert "První řádek s popisky" in html


def test_delimiter_detection_skips_leading_blank_lines():
    content = b"\n" * 6 + b"Name;ISIN\nAlfa a.s.;CZ0005112300\n"
    assert _entities(content) == [("Alfa a.s.", "CZ0005112300")]


# Size, speed and memory.

def test_too_many_entities_stops_early_and_says_more_than():
    english, czech = _refusal(_csv([GOOD_ROW] * 150))
    assert english == (
        f"{TOO_MANY_EN} The maximum is {MAX_ENTITIES} per file."
    )
    assert czech == f"{TOO_MANY_CS} Maximum je {MAX_ENTITIES} na soubor."

    started = time.perf_counter()
    english, _ = _refusal(b"A\n" * 2_000_000, "big.txt")
    assert english.startswith(TOO_MANY_EN)
    assert time.perf_counter() - started < 10


def test_names_among_megabytes_of_blank_lines_are_read_quickly():
    names = b"".join(b"Firma %d\n" % index for index in range(100))
    content = names + b"\n" * (4 * 1024 * 1024 - 3000)
    started = time.perf_counter()
    assert len(parse_upload("in.txt", content)) == 100
    assert time.perf_counter() - started < 30


def test_xlsx_rows_past_a_stale_dimension_are_read():
    rows_xml = "".join(
        _inline_row(number, f"Company {number}", "CZ0005112300")
        for number in range(1, 6)
    )
    content = _with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml, dimension="A1"),
    })
    assert _entities(content, "in.xlsx") == [
        (f"Company {number}", "CZ0005112300") for number in range(1, 6)
    ]


def test_one_far_away_xlsx_cell_is_read_quickly():
    workbook = Workbook()
    workbook.active.append(["Tiny company", "CZ0005112300"])
    workbook.active["XFD1048576"] = "x"
    content = io.BytesIO()
    workbook.save(content)
    started = time.perf_counter()
    assert _entities(content.getvalue(), "in.xlsx") == [
        ("Tiny company", "CZ0005112300"),
    ]
    assert time.perf_counter() - started < 20


def test_a_long_xlsx_stops_at_the_cap():
    rows_xml = "".join(
        _inline_row(number, f"Company {number}")
        for number in range(1, 100_001)
    )
    content = _with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml, "A1:A100000"),
    })
    started = time.perf_counter()
    english, czech = _refusal(content, "in.xlsx")
    assert english.startswith(TOO_MANY_EN)
    assert czech.startswith(TOO_MANY_CS)
    assert time.perf_counter() - started < 15


def test_many_semicolon_lines_in_xlsx_are_too_many():
    rows = [[f"Firma {index};CZ0005112300"] for index in range(150)]
    english, _ = _refusal(_xlsx(rows), "in.xlsx")
    assert english.startswith(TOO_MANY_EN)


def test_an_xlsx_row_that_is_not_a_semicolon_line_keeps_columns_as_is():
    rows = [["Alfa;Beta"], ["Plain Company", "CZ0005112300"]]
    assert _entities(_xlsx(rows), "in.xlsx") == [
        ("Alfa;Beta", None), ("Plain Company", "CZ0005112300"),
    ]


def test_a_long_xlsx_of_semicolon_lines_stops_at_the_cap(monkeypatch):
    read = []
    filled_rows = upload._filled_rows

    def counted_rows(sheet, width):
        for row in filled_rows(sheet, width):
            read.append(row)
            yield row

    monkeypatch.setattr(upload, "_filled_rows", counted_rows)
    # Each line was split again at its commas, as Excel does.
    rows_xml = "".join(
        _inline_row(
            number, f"Firma {number};CZ0005112300;CZ;Praha;Ulice 1",
            " 2. patro", " vchod B", ";110 00",
        )
        for number in range(1, 5001)
    )
    content = _with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml, "A1:D5000"),
    })
    english, _ = _refusal(content, "in.xlsx")
    assert english.startswith(TOO_MANY_EN)
    assert len(read) < 3 * MAX_ENTITIES


def _assert_too_much_data_quickly(content):
    started = time.perf_counter()
    english, czech = _refusal(content, "in.xlsx")
    assert english.startswith(TOO_MUCH_DATA_EN)
    assert czech.startswith(TOO_MUCH_DATA_CS)
    assert time.perf_counter() - started < 30


def test_a_row_of_countless_empty_cells_is_refused_quickly():
    # openpyxl builds every cell of a row before any column limit.
    rows_xml = (
        _inline_row(1, "Tiny company") + "<row>" + "<c/>" * 600_000
        + "</row>"
    )
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml, "A1:B2"),
    }))


def test_countless_empty_rows_are_refused_quickly():
    rows_xml = _inline_row(1, "Tiny company") + "<row/>" * 600_000
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml, "A1:B2"),
    }))


def _listed_sheets(count):
    """Rewrite workbook.xml to list its one sheet ``count`` times."""
    def rewrite(workbook):
        entries = "".join(
            f'<sheet name="S{index}" sheetId="{index + 1}" r:id="rId1" />'
            for index in range(count)
        )
        start = workbook.index(b"<sheets>") + len(b"<sheets>")
        end = workbook.index(b"</sheets>")
        return workbook[:start] + entries.encode() + workbook[end:]
    return rewrite


def test_one_sheet_listed_many_times_is_refused_quickly():
    # Loading reads the part of every sheet entry; with no <dimension>
    # it reads all of it, every time.
    rows_xml = _inline_row(1, "Tiny company") + "<row/>" * 20_000
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml),
        "xl/workbook.xml": _listed_sheets(40),
    }))


def test_one_big_part_read_many_times_is_refused_quickly():
    rows_xml = _inline_row(1, "Tiny company") + " " * (5 * 1024 * 1024)
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml),
        "xl/workbook.xml": _listed_sheets(40),
    }))


def test_shared_strings_under_any_part_name_are_bounded():
    # openpyxl finds the shared strings through [Content_Types].xml,
    # whatever the part is called.
    strings = (
        b'<?xml version="1.0"?><sst xmlns="http://schemas.openxmlformats'
        b'.org/spreadsheetml/2006/main">' + b"<si><t>A</t></si>" * 300_000
        + b"</sst>"
    )
    override = (
        b'<Override PartName="/xl/strings.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings'
        b'+xml" /></Types>'
    )
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/strings.xml": strings,
        "[Content_Types].xml": lambda old: old.replace(b"</Types>", override),
    }))


def test_one_chart_drawn_many_times_is_refused_quickly():
    workbook = Workbook()
    workbook.active.append(["Tiny company", 1])
    chart = BarChart()
    chart.add_data(Reference(workbook.active, min_col=2, min_row=1))
    workbook.create_chartsheet().add_chart(chart)
    content = io.BytesIO()
    workbook.save(content)

    def drawing(old):
        start = old.index(b"<absoluteAnchor>")
        end = old.index(b"</absoluteAnchor>") + len(b"</absoluteAnchor>")
        return old[:start] + old[start:end] * 500 + old[end:]

    def chart_sheet_rels(old):
        start = old.index(b"<Relationship ")
        end = old.index(b"/>", start) + 2
        relationships = b"".join(
            old[start:end].replace(b'Id="rId1"', b'Id="rId%d"' % number)
            for number in range(1, 11)
        )
        return old[:start] + relationships + old[end:]

    _assert_too_much_data_quickly(_with_parts(content.getvalue(), {
        "xl/drawings/drawing1.xml": drawing,
        "xl/chartsheets/_rels/sheet1.xml.rels": chart_sheet_rels,
    }))


def test_an_xlsx_part_that_unpacks_to_too_much_is_refused():
    sheet = _sheet_xml(_inline_row(1, "Tiny company")).replace(
        b"</worksheet>", b" " * (60 * 1024 * 1024) + b"</worksheet>"
    )
    _assert_too_much_data_quickly(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": sheet,
    }))


def _assert_unreadable(content):
    english, czech = _refusal(content, "in.xlsx")
    assert english == (
        "Could not read the .xlsx file; is it a valid Excel file?"
    )
    assert czech.startswith("Soubor .xlsx se nepodařilo přečíst.")


def test_a_row_past_the_last_excel_row_is_refused():
    # openpyxl would first yield every empty row before it.
    rows_xml = (
        _inline_row(1, "Tiny company")
        + _inline_row(20_000_000, "Late company")
    )
    _assert_unreadable(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": _sheet_xml(rows_xml),
    }))


def test_an_xlsx_part_with_a_dtd_is_refused():
    # Office Open XML has no DTDs, and their entities can blow a small
    # part up to gigabytes of text.
    sheet = _sheet_xml(_inline_row(1, "&company;")).replace(
        b"<worksheet",
        b'<!DOCTYPE worksheet [<!ENTITY company "Tiny company">]>'
        b"<worksheet",
        1,
    )
    _assert_unreadable(_with_parts(_xlsx([["x"]]), {
        "xl/worksheets/sheet1.xml": sheet,
    }))


def test_a_file_that_is_not_a_workbook_gets_the_clean_message():
    _assert_unreadable(b"not a zip at all")

