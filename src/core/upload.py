
"""Parse an uploaded .xlsx/.csv file into a list of InputEntity objects.

Columns are read by position, in the order the bulk form documents:
Name, ISIN, Country, City, Street, Postal code. A first row that looks
like a header is skipped. Uses openpyxl for .xlsx and the stdlib csv
module for .csv, so no extra dependency is needed.
"""

import csv
import io
import logging
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from .models import InputEntity

logger = logging.getLogger(__name__)

#: Maximum number of entities accepted in one bulk upload.
MAX_ENTITIES = 100

#: Columns in the documented order: position -> InputEntity field.
_COLUMNS = ("name", "isin", "country", "town", "street", "zip_code")

#: First-cell values (lower-cased) that mark a header row to skip.
_HEADER_CELLS = frozenset({
    "name", "entity", "entity name", "company", "company name",
    "issuer", "issuer name", "firma", "nazev", "název",
})


def parse_upload(filename: str, content: bytes) -> list[InputEntity]:
    """Parse uploaded bytes into entities.

    Args:
        filename: The original file name (its extension picks the parser).
        content: The raw file bytes.

    Returns:
        The parsed entities (rows with a name or an ISIN).

    Raises:
        ValueError: For an empty/unreadable file, an unsupported
            extension, no usable rows, or more than MAX_ENTITIES rows.
    """
    if not content:
        raise ValueError("The file is empty.")

    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        rows = _read_xlsx(content)
    elif ext == ".csv":
        rows = _read_csv(content)
    else:
        raise ValueError(
            "Unsupported file type. Please upload a .xlsx or .csv file."
        )

    rows = _drop_header(rows)
    entities = _rows_to_entities(rows)
    if not entities:
        raise ValueError(
            "No entities found. Each row needs a name (first column) or "
            "an ISIN (second column)."
        )
    if len(entities) > MAX_ENTITIES:
        raise ValueError(
            f"Too many entities ({len(entities)}). The maximum is "
            f"{MAX_ENTITIES} per file."
        )
    return entities


def _read_xlsx(content: bytes) -> list[list[str]]:
    """Read the first worksheet into rows of trimmed string cells."""
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
    except Exception as error:
        logger.warning("Failed to parse uploaded .xlsx file: %s", error)
        raise ValueError(
            "Could not read the .xlsx file; is it a valid Excel file?"
        ) from error

    sheet = workbook.active
    rows = []
    for raw in sheet.iter_rows(values_only=True):
        rows.append(
            ["" if cell is None else str(cell).strip() for cell in raw]
        )
    workbook.close()
    return rows


def _decode(content: bytes) -> str:
    """Decode CSV bytes, trying common encodings (incl. Czech cp1250)."""
    for encoding in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return content.decode("latin-1")  # latin-1 never fails


def _read_csv(content: bytes) -> list[list[str]]:
    """Read CSV bytes into rows, auto-detecting comma vs semicolon."""
    text = _decode(content)
    sample = "\n".join(text.splitlines()[:5])
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell.strip() for cell in row] for row in reader]


def _drop_header(rows: list[list[str]]) -> list[list[str]]:
    """Drop the first row when it looks like a header."""
    if not rows or not rows[0]:
        return rows
    first = rows[0][0].strip().lower()
    second = rows[0][1].strip().lower() if len(rows[0]) > 1 else ""
    if first in _HEADER_CELLS or second == "isin":
        return rows[1:]
    return rows


def _rows_to_entities(rows: list[list[str]]) -> list[InputEntity]:
    """Map positional columns to entities, skipping empty/invalid rows."""
    entities: list[InputEntity] = []
    for row in rows:
        values = {}
        for index, field in enumerate(_COLUMNS):
            cell = row[index].strip() if index < len(row) else ""
            values[field] = cell or None
        # Keep a row that has a name or an ISIN; skip only when both are
        # missing. InputEntity enforces the same name-or-ISIN rule.
        if not values["name"] and not values["isin"]:
            continue
        try:
            entities.append(InputEntity(**values))
        except ValidationError:
            continue  # skip a row that fails validation (e.g. oversized)
    return entities

