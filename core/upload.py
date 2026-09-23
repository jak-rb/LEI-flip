
"""Parse an uploaded .xlsx, .csv, .tsv or .txt file into entities.

Columns are read by position, in the order the bulk form documents:
Name, ISIN, Country, City, Street, Postal code. The first non-blank
row is skipped when it holds column labels (a header). Uses openpyxl
for .xlsx and the stdlib csv module for the text formats, so no extra
dependency is needed.

A row with a name or an ISIN is never dropped or cut short: one with a
value over its length limit refuses the whole file, with a message
naming the row. Reading stops as soon as a file is known to hold too
many entities.
"""

import codecs
import csv
import io
import itertools
import logging
import re
import zipfile
from collections.abc import Iterable, Iterator
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chartsheet import Chartsheet
from pydantic import ValidationError
from unidecode import unidecode

from .isin import is_valid_isin, normalize_isin
from .models import InputEntity, InputError

logger = logging.getLogger(__name__)

#: Maximum number of entities accepted in one bulk upload.
MAX_ENTITIES = 100

#: Columns in the documented order: position -> InputEntity field.
_COLUMNS = ("name", "isin", "country", "town", "street", "zip_code")

#: A row as read from a file: its number as the user sees it there
#: (1-based, counting the header and blank rows) and its trimmed cells.
_Row = tuple[int, list[str]]

#: How a message names each column: field -> (English, Czech).
_FIELD_NAMES = {
    "name": ("the name", "název"),
    "isin": ("the ISIN", "ISIN"),
    "country": ("the country", "země"),
    "town": ("the city", "město"),
    "street": ("the street", "ulice"),
    "zip_code": ("the postal code", "PSČ"),
}

#: At most this many rows are named in one message.
_LISTED_ROWS = 5

#: Words the column labels of a header are made of, lower-case and
#: without diacritics: "Název subjektu", "Legal name",
#: "PARTY_FULL_NAME", "ISIN kód", "ISIN_IDENT", "Země" and "PSČ" are
#: all labels. A cell with any other word, such as a company name, is
#: data.
_HEADER_WORDS = frozenset({
    # Name
    "name", "entity", "company", "legal", "full", "party", "issuer",
    "firm", "firma", "firmy", "nazev", "subjekt", "subjektu",
    "obchodni", "jmeno", "emitent", "emitenta", "spolecnost",
    "spolecnosti",
    # ISIN
    "isin", "code", "kod", "ident", "identifier",
    # Address
    "country", "zeme", "stat", "city", "town", "mesto", "obec",
    "street", "ulice", "address", "adresa", "addr", "zip", "postal",
    "postcode", "psc",
})

#: Columns read from an .xlsx row whose column A holds a semicolon
#: line: Excel split the line again at every comma, so rebuilding it
#: takes the cells after the six documented columns too.
_LINE_COLUMNS = 50

#: Largest unpacked size of an .xlsx, in all and of its shared strings,
#: that is read. A real 100-entity workbook is far smaller, while a
#: crafted file of a few kilobytes can unpack to gigabytes, or hold
#: millions of tiny shared strings that take minutes to load (about
#: 10 microseconds each).
_MAX_UNPACKED_BYTES = 50 * 1024 * 1024
_MAX_SHARED_STRINGS_BYTES = 10 * 1024 * 1024

#: Text formats read with the csv module: extension -> delimiter, where
#: None means detect it. Excel's "Text (Tab delimited)" and "Unicode
#: Text" exports, and cells copied from Excel into a text editor, are
#: tab-separated .txt; a .tsv is tab-separated by definition.
_TEXT_DELIMITERS = {".csv": None, ".tsv": "\t", ".txt": "\t"}

#: Shown (English, Czech) for a text file that cannot be read as rows.
_UNREADABLE_TEXT = (
    "Could not read the file. Save it from Excel as CSV UTF-8, or "
    "upload the .xlsx instead.",
    "Soubor se nepodařilo přečíst. Uložte ho z Excelu jako CSV UTF-8 "
    "nebo nahrajte soubor .xlsx.",
)


def parse_upload(filename: str, content: bytes) -> list[InputEntity]:
    """Parse uploaded bytes into entities.

    Args:
        filename: The original file name (its extension picks the parser).
        content: The raw file bytes.

    Returns:
        The parsed entities (rows with a name or an ISIN).

    Raises:
        InputError: For an empty/unreadable file, an unsupported
            extension, no usable rows, more than MAX_ENTITIES rows, or
            rows with a value over its length limit.
    """
    if not content:
        raise InputError("The file is empty.", "Soubor je prázdný.")

    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        rows = _read_xlsx(content)
    elif ext in _TEXT_DELIMITERS:
        rows = _read_csv(content, _TEXT_DELIMITERS[ext])
    else:
        raise InputError(
            "Unsupported file type. Please upload a .xlsx, .csv, .tsv or "
            ".txt file.",
            "Nepodporovaný typ souboru. Nahrajte prosím soubor .xlsx, "
            ".csv, .tsv nebo .txt.",
        )

    entities = _rows_to_entities(rows)
    if not entities:
        raise InputError(
            "No entities found. Each row needs a name (first column) or "
            "an ISIN (second column).",
            "Nebyly nalezeny žádné subjekty. Každý řádek musí obsahovat "
            "název (první sloupec) nebo ISIN (druhý sloupec).",
        )
    return entities


def _read_xlsx(content: bytes) -> list[_Row]:
    """Read the active worksheet's entity rows (see _entity_rows)."""
    # The whole read sits in the try: in read-only mode a damaged sheet
    # only fails once its rows are read.
    try:
        _check_unpacked_size(content)
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        sheet = workbook.active
        # Excel saves the sheet on screen as the active one; a chart
        # sheet has no cells, so fall back to the first worksheet.
        if isinstance(sheet, Chartsheet):
            sheet = workbook.worksheets[0]
        # Read-only mode stops at the size the sheet states, which some
        # writers leave stale: the rows past it would be lost.
        sheet.reset_dimensions()
        if _holds_semicolon_lines(sheet):
            rows = _entity_rows(_split_semicolon_lines(sheet))
        else:
            rows = _entity_rows(
                (number, [cell.strip() for cell in cells])
                for number, cells in _filled_rows(sheet, len(_COLUMNS))
            )
        workbook.close()
    except InputError:
        raise
    except Exception as error:
        logger.warning("Failed to parse uploaded .xlsx file: %s", error)
        raise InputError(
            "Could not read the .xlsx file; is it a valid Excel file?",
            "Soubor .xlsx se nepodařilo přečíst. Je to platný soubor "
            "Excelu?",
        ) from error
    return rows


def _check_unpacked_size(content: bytes) -> None:
    """Refuse an .xlsx that unpacks to more than is safe to read."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
    # zipfile never unpacks a member past its stated size, so the
    # stated sizes bound what openpyxl can be made to parse.
    total = sum(member.file_size for member in members)
    shared_strings = sum(
        member.file_size for member in members
        if member.filename.lower().endswith("sharedstrings.xml")
    )
    if (
        total > _MAX_UNPACKED_BYTES
        or shared_strings > _MAX_SHARED_STRINGS_BYTES
    ):
        logger.warning(
            "Uploaded .xlsx unpacks to %d bytes (%d of shared strings)",
            total, shared_strings,
        )
        raise InputError(
            "The .xlsx file holds too much data to read. Keep only the "
            "sheet with the entities, or save it as CSV UTF-8.",
            "Soubor .xlsx obsahuje příliš mnoho dat. Ponechte v něm jen "
            "list se subjekty nebo ho uložte jako CSV UTF-8.",
        )


def _filled_rows(sheet, width: int) -> Iterator[tuple[int, list[str]]]:
    """Yield (worksheet row number, cells as text) of non-blank rows.

    Reads the first ``width`` columns of every row.
    """
    rows = sheet.iter_rows(max_col=width, values_only=True)
    for number, values in enumerate(rows, start=1):
        # Checked on the raw values first, as it is cheap: one stray
        # far-away cell makes openpyxl yield a million empty rows.
        if values.count(None) == len(values):
            continue
        cells = ["" if cell is None else str(cell) for cell in values]
        if any(cell.strip() for cell in cells):
            yield number, cells


def _holds_semicolon_lines(sheet) -> bool:
    """Whether every non-blank row keeps a semicolon line in column A.

    That is what Excel shows for a semicolon CSV opened with the comma
    as the delimiter: each whole line lands in column A, split again
    into the next columns wherever a value holds a comma.
    """
    found = False
    for _, cells in _filled_rows(sheet, _LINE_COLUMNS):
        if ";" not in cells[0]:
            return False
        found = True
    return found


def _split_semicolon_lines(sheet) -> Iterator[_Row]:
    """Rebuild each row's original line and split it at semicolons."""
    for number, cells in _filled_rows(sheet, _LINE_COLUMNS):
        while not cells[-1].strip():
            cells.pop()
        # Excel split the line at its commas, so joining the cells with
        # commas restores it (a cell keeps its leading space). Each line
        # is parsed on its own: an unbalanced quote must not swallow
        # the rows after it.
        line = ",".join(cells)
        fields = next(csv.reader([line], delimiter=";"), [])
        yield number, [field.strip() for field in fields]


def _decode(content: bytes) -> str:
    """Decode CSV bytes, trying common encodings (incl. Czech cp1250)."""
    encodings = ("utf-8-sig", "utf-8", "cp1250", "latin-1")
    # UTF-16 (what Windows PowerShell's `>` writes, for one) starts
    # with a byte-order mark and would decode as garbage below.
    if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        encodings = ("utf-16",) + encodings
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, ValueError):
            continue
    return content.decode("latin-1")  # latin-1 never fails


def _read_csv(content: bytes, delimiter: str | None) -> list[_Row]:
    """Read delimited text's entity rows; a None delimiter is detected.

    Detection picks whichever of comma, semicolon and tab is most common
    in the first non-blank lines; a tie keeps the comma, then the
    semicolon.
    """
    text = _decode(content)
    # Text never holds a NUL: this is binary content, such as an Excel
    # workbook renamed to .csv, which would only parse into garbage.
    if "\x00" in text:
        logger.warning("Uploaded text file is binary (contains NUL)")
        raise InputError(*_UNREADABLE_TEXT)
    if delimiter is None:
        lines = io.StringIO(text, newline="")
        sample = "".join(
            itertools.islice((line for line in lines if line.strip()), 5)
        )
        delimiter = max((",", ";", "\t"), key=sample.count)
    # newline="" leaves line endings to the csv module, which accepts
    # \n, \r\n and the lone \r of old Mac files alike.
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    rows = (
        (number, [cell.strip() for cell in row[:len(_COLUMNS)]])
        for number, row in enumerate(reader, start=1)
        if row
    )
    try:
        return _entity_rows(rows)
    except csv.Error as error:  # e.g. a cell over the 128 KB limit
        logger.warning("Failed to parse uploaded text file: %s", error)
        raise InputError(*_UNREADABLE_TEXT) from error


def _entity_rows(rows: Iterable[_Row]) -> list[_Row]:
    """Keep the rows that hold a name or an ISIN, cut to the columns.

    Blank rows are skipped, and so is the first non-blank row when it
    is a header. Reading stops at the first row over MAX_ENTITIES, so a
    huge file is refused without being read to its end.

    Args:
        rows: The file's rows, in order.

    Returns:
        The rows to build entities from.

    Raises:
        InputError: When more than MAX_ENTITIES rows hold an entity.
    """
    kept = []
    header_checked = False
    for number, cells in rows:
        cells = cells[:len(_COLUMNS)]
        if not any(cells):
            continue
        if not header_checked:
            header_checked = True
            if _is_header(cells):
                continue
        # Counted before validation: an invalid row counts toward the
        # cap too, and "too many" is then the error to show.
        if not any(cells[:2]):
            continue
        kept.append((number, cells))
        if len(kept) > MAX_ENTITIES:
            raise InputError(
                f"Too many entities (more than {MAX_ENTITIES}). The "
                f"maximum is {MAX_ENTITIES} per file.",
                f"Příliš mnoho subjektů (více než {MAX_ENTITIES}). "
                f"Maximum je {MAX_ENTITIES} na soubor.",
            )
    return kept


def _is_header(cells: list[str]) -> bool:
    """Whether a row holds column labels rather than an entity."""
    # A company may be named like a label, but a real ISIN beside the
    # name shows the row is data.
    if len(cells) > 1 and is_valid_isin(normalize_isin(cells[1])):
        return False
    return any(_is_label(cell) for cell in cells)


def _is_label(cell: str) -> bool:
    """Whether a cell is a column label, such as "ISIN kód"."""
    words = re.findall(r"[a-z0-9]+", unidecode(cell).lower())
    return bool(words) and all(word in _HEADER_WORDS for word in words)


def _rows_to_entities(rows: list[_Row]) -> list[InputEntity]:
    """Map positional columns to one entity per row.

    Raises:
        InputError: Naming the rows with a value over its length limit
            (precision first: a value is never cut short to fit).
    """
    entities: list[InputEntity] = []
    too_long: list[tuple[int, list[tuple[str, int]]]] = []
    for number, row in rows:
        values = {}
        for index, field in enumerate(_COLUMNS):
            cell = row[index].strip() if index < len(row) else ""
            values[field] = cell or None
        try:
            entities.append(InputEntity(**values))
        except ValidationError as error:
            # Each row has a name or an ISIN and only text values, so a
            # length limit is the one check it can fail.
            limits = {
                issue["loc"][0]: issue["ctx"]["max_length"]
                for issue in error.errors()
            }
            too_long.append((number, [
                (field, limits[field]) for field in _COLUMNS
                if field in limits
            ]))
    if too_long:
        raise _too_long_error(too_long)
    return entities


def _too_long_error(
    rows: list[tuple[int, list[tuple[str, int]]]]
) -> InputError:
    """The error naming the rows that have a value over its limit."""
    english, czech = [], []
    for number, limits in rows[:_LISTED_ROWS]:
        english.append(f"Row {number}: " + "; ".join(
            f"{_FIELD_NAMES[field][0]} is longer than {limit} characters"
            for field, limit in limits
        ) + ".")
        czech.append(f"Řádek {number}: " + "; ".join(
            f"{_FIELD_NAMES[field][1]} je delší než {limit} znaků"
            for field, limit in limits
        ) + ".")
    more = len(rows) - _LISTED_ROWS
    if more > 0:
        english.append(f"And {more} more row{'s' if more > 1 else ''}.")
        if more == 1:
            more_cs = "další řádek"
        elif more < 5:
            more_cs = "další řádky"
        else:
            more_cs = "dalších řádků"
        czech.append(f"A ještě {more} {more_cs}.")
    return InputError(" ".join(english), " ".join(czech))

